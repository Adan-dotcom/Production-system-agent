"""
Aspel Excel parser — reads and normalizes the Excel file exported from Aspel SAE.

Key rules:
- CVE_DOC    → sae_order_number  (one order per unique CVE_DOC)
- CVE_ART    → product_code      (one line per CVE_DOC+CVE_ART)
- CANT       → quantity_ordered  (commercial qty — NOT kg, never assume)
- KILOSxPARTIDA → target_kg     (nullable — missing = not blocking)
- Idempotent: run N times, no duplicates
- Logs errors per-row without stopping
- Detects header row automatically
"""

import json
import re
from datetime import datetime, date
from typing import Any

import openpyxl

# ── Column aliases accepted from Aspel Excel ─────────────────────────────────
COLUMN_MAP = {
    # canonical_name : list of possible header strings (upper, stripped)
    "cve_clpv":        ["CVE_CLPV", "CLAVE CLIENTE", "CVE CLIENTE"],
    "nombre_cliente":  ["NOMBRE CLIENTE", "CLIENTE", "NOMBRE_CLIENTE"],
    "cve_vend":        ["CVE_VEND", "CLAVE VENDEDOR"],
    "nombre_vend":     ["NOMBRE", "VENDEDOR", "NOMBRE VENDEDOR"],
    "cve_art":         ["CVE_ART", "CLAVE ARTICULO", "ARTICULO", "CLAVE ART"],
    "descr":           ["DESCR", "DESCRIPCION", "DESCRIPCIÓN", "DESCRIPTION"],
    "cve_doc":         ["CVE_DOC", "PEDIDO", "FOLIO", "DOCUMENTO", "NUM_DOC"],
    "status":          ["STATUS", "ESTADO", "ESTATUS"],
    "fecha_doc":       ["FECHA_DOC", "FECHA DOC", "FECHA PEDIDO", "F_DOC"],
    "fecha_ent":       ["FECHA_ENT", "FECHA ENT", "FECHA ENTREGA", "F_ENT"],
    "cant":            ["CANT", "CANTIDAD", "QTY"],
    "prec":            ["PREC", "PRECIO", "PRICE"],
    "tot_partida":     ["TOT_PARTIDA", "TOTAL PARTIDA", "TOTAL", "IMPORTE"],
    "doc_sig":         ["DOC_SIG", "DOC SIG", "DOCUMENTO SIGUIENTE"],
    "calle":           ["CALLE", "DIRECCION", "DIRECCIÓN", "DOMICILIO"],
    "kilos_partida":   ["KILOSxPARTIDA", "KILOS X PARTIDA", "KILOS_PARTIDA",
                        "KILOS PARTIDA", "KG_PARTIDA", "KGS"],
}

REQUIRED_COLS = {"cve_doc", "cve_art", "nombre_cliente"}


# ── Helpers ───────────────────────────────────────────────────────────────────

def _clean_str(val: Any) -> str:
    """Strip, collapse whitespace, remove ## artifacts."""
    if val is None:
        return ""
    s = str(val).strip()
    # Remove ## delimiters — take first clean token
    if "##" in s:
        tokens = [t.strip() for t in s.split("##") if t.strip()]
        s = tokens[0] if tokens else ""
    # Collapse multiple spaces
    s = re.sub(r"\s{2,}", " ", s)
    return s


def _clean_code(val: Any) -> str:
    """Like _clean_str but also removes internal spaces (codes like 'CVE 001' → 'CVE001')."""
    s = _clean_str(val)
    return re.sub(r"\s+", "", s)


def _excel_date(val: Any) -> datetime | None:
    """Convert Excel serial number or string to datetime."""
    if val is None or (isinstance(val, str) and not val.strip()):
        return None
    if isinstance(val, (datetime, date)):
        return datetime.combine(val, datetime.min.time()) if isinstance(val, date) else val
    if isinstance(val, (int, float)):
        # Excel serial: days since 1900-01-00
        try:
            from openpyxl.utils.datetime import from_excel
            return from_excel(val)
        except Exception:
            return None
    # Try common string formats
    s = str(val).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y", "%Y%m%d"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def _safe_float(val: Any) -> float | None:
    if val is None or (isinstance(val, str) and not val.strip()):
        return None
    try:
        return float(str(val).replace(",", "").strip())
    except (ValueError, TypeError):
        return None


# ── Header detection ──────────────────────────────────────────────────────────

def _extract_measure_caliber(product_code: str, description: str = "") -> tuple[str | None, str | None]:
    """
    Extract measure/caliber from Aspel product text.
    Measure is before '#'; caliber is the first numeric token after '#'.
    """
    source = product_code or description or ""
    if "#" not in source and description and "#" in description:
        source = description
    if "#" not in source:
        return None, None

    before, after = source.split("#", 1)
    measure = re.sub(r"[^0-9.+xX]", "", before).upper()
    measure = re.sub(r"X+", "X", measure).strip(".+X")

    caliber_match = re.search(r"\d+(?:\.\d+)?", after)
    caliber = caliber_match.group(0) if caliber_match else None

    return measure or None, caliber


def _detect_header_row(ws) -> tuple[int, dict[str, int]]:
    """
    Scan the worksheet to find the row containing column headers.
    Returns (header_row_index_1based, {canonical_name: col_index_0based}).
    Raises ValueError if headers not found.
    """
    # Build reverse map: UPPER_HEADER_STRING → canonical_name
    reverse: dict[str, str] = {}
    for canonical, aliases in COLUMN_MAP.items():
        for alias in aliases:
            reverse[alias.upper()] = canonical

    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_idx > 30:
            break  # give up after 30 rows
        found: dict[str, int] = {}
        for col_idx, cell_val in enumerate(row):
            if cell_val is None:
                continue
            key = str(cell_val).strip().upper()
            if key in reverse:
                canonical = reverse[key]
                if canonical not in found:
                    found[canonical] = col_idx

        # Require at least cve_doc and (cve_art or descr) and nombre_cliente
        has_doc = "cve_doc" in found
        has_art = "cve_art" in found or "descr" in found
        has_cli = "nombre_cliente" in found
        if has_doc and has_art and has_cli:
            return row_idx, found

    raise ValueError(
        "No se encontraron los encabezados requeridos (CVE_DOC, CVE_ART/DESCR, NOMBRE CLIENTE) "
        "en las primeras 30 filas del Excel."
    )


# ── Main parse function ───────────────────────────────────────────────────────

def parse_aspel_excel(file_bytes: bytes) -> dict:
    """
    Parse an Aspel Excel file.

    Returns:
    {
        "orders": {
            sae_order_number: {
                "sae_order_number": str,
                "customer_name": str,
                "customer_code": str,
                "seller_code": str,
                "seller_name": str,
                "order_status": str,
                "order_date": datetime|None,
                "delivery_date": datetime|None,
                "shipping_address": str,
                "target_kg": float|None,   # sum of KILOSxPARTIDA for this order
                "lines": [
                    {
                        "product_code": str,
                        "product_description": str,
                        "quantity_ordered": float|None,
                        "target_kg": float|None,
                        "unit_price": float|None,
                        "line_total": float|None,
                        "following_document": str,
                        "line_number": int,   # sequential within order
                    }, ...
                ],
            }, ...
        },
        "raw_rows": [
            {
                "row_number": int,
                "raw": dict,
                "normalized": dict,
                "status": "ok"|"warning"|"error"|"skipped",
                "errors": str,
            }, ...
        ],
        "header_row": int,
        "total_rows": int,
    }
    """
    import io
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    ws = wb.active

    header_row_idx, col_map = _detect_header_row(ws)

    def get(row_values: tuple, canonical: str, default=None):
        idx = col_map.get(canonical)
        if idx is None or idx >= len(row_values):
            return default
        return row_values[idx]

    orders: dict[str, dict] = {}
    raw_rows: list[dict] = []
    total_data_rows = 0

    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_idx <= header_row_idx:
            continue

        # Skip completely empty rows
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in row):
            continue

        total_data_rows += 1
        absolute_row = row_idx

        # Build raw snapshot
        raw = {}
        for canonical, col_idx in col_map.items():
            if col_idx < len(row):
                raw[canonical] = str(row[col_idx]) if row[col_idx] is not None else None

        errors = []
        status = "ok"

        # Normalize fields
        cve_doc        = _clean_code(get(row, "cve_doc"))
        nombre_cliente = _clean_str(get(row, "nombre_cliente"))
        cve_art        = _clean_code(get(row, "cve_art"))
        descr          = _clean_str(get(row, "descr"))
        cve_clpv       = _clean_code(get(row, "cve_clpv"))
        cve_vend       = _clean_code(get(row, "cve_vend"))
        nombre_vend    = _clean_str(get(row, "nombre_vend"))
        order_status   = _clean_str(get(row, "status"))
        fecha_doc      = _excel_date(get(row, "fecha_doc"))
        fecha_ent      = _excel_date(get(row, "fecha_ent"))
        cant           = _safe_float(get(row, "cant"))
        prec           = _safe_float(get(row, "prec"))
        tot_partida    = _safe_float(get(row, "tot_partida"))
        doc_sig        = _clean_str(get(row, "doc_sig"))
        calle          = _clean_str(get(row, "calle"))
        kilos          = _safe_float(get(row, "kilos_partida"))
        measure, caliber = _extract_measure_caliber(cve_art, descr)

        # Validation
        if not cve_doc:
            errors.append("CVE_DOC vacío — fila ignorada")
            status = "error"
        if not nombre_cliente:
            errors.append("NOMBRE CLIENTE vacío")
            status = "warning" if status == "ok" else status
        if not cve_art and not descr:
            errors.append("CVE_ART y DESCR ambos vacíos")
            status = "warning" if status == "ok" else status

        normalized = {
            "cve_doc":        cve_doc,
            "nombre_cliente": nombre_cliente,
            "cve_clpv":       cve_clpv,
            "cve_vend":       cve_vend,
            "nombre_vend":    nombre_vend,
            "cve_art":        cve_art,
            "descr":          descr,
            "order_status":   order_status,
            "fecha_doc":      fecha_doc.isoformat() if fecha_doc else None,
            "fecha_ent":      fecha_ent.isoformat() if fecha_ent else None,
            "cant":           cant,
            "prec":           prec,
            "tot_partida":    tot_partida,
            "doc_sig":        doc_sig,
            "calle":          calle,
            "kilos_partida":  kilos,
            "measure":        measure,
            "caliber":        caliber,
        }

        raw_rows.append({
            "row_number": absolute_row,
            "raw":        raw,
            "normalized": normalized,
            "status":     "skipped" if status == "error" else status,
            "errors":     "; ".join(errors) if errors else "",
        })

        if status == "error":
            continue

        # ── Group into orders dict ─────────────────────────────────────────
        if cve_doc not in orders:
            orders[cve_doc] = {
                "sae_order_number":  cve_doc,
                "customer_name":     nombre_cliente,
                "customer_code":     cve_clpv,
                "seller_code":       cve_vend,
                "seller_name":       nombre_vend,
                "order_status":      order_status,
                "order_date":        fecha_doc,
                "delivery_date":     fecha_ent,
                "shipping_address":  calle,
                "target_kg":         None,
                "lines":             [],
            }

        order = orders[cve_doc]

        # Update order-level fields from later rows (more data = better)
        if not order["customer_name"] and nombre_cliente:
            order["customer_name"] = nombre_cliente
        if not order["seller_name"] and nombre_vend:
            order["seller_name"] = nombre_vend
        if not order["order_date"] and fecha_doc:
            order["order_date"] = fecha_doc
        if not order["delivery_date"] and fecha_ent:
            order["delivery_date"] = fecha_ent
        if not order["shipping_address"] and calle:
            order["shipping_address"] = calle

        # Accumulate target_kg
        if kilos is not None:
            order["target_kg"] = (order["target_kg"] or 0.0) + kilos

        # Add line — deduplicate by product_code within the order
        existing_line = next(
            (l for l in order["lines"]
             if l["product_code"] == cve_art and cve_art),
            None
        )
        if existing_line:
            # Update with more recent values if better
            if kilos is not None:
                existing_line["target_kg"] = kilos
            if cant is not None:
                existing_line["quantity_ordered"] = cant
            if measure and not existing_line.get("measure"):
                existing_line["measure"] = measure
            if caliber and not existing_line.get("caliber"):
                existing_line["caliber"] = caliber
        else:
            line_number = len(order["lines"]) + 1
            order["lines"].append({
                "product_code":        cve_art,
                "product_description": descr or cve_art,
                "measure":             measure,
                "caliber":             caliber,
                "quantity_ordered":    cant,    # commercial qty — NOT kg
                "target_kg":           kilos,   # nullable — not required
                "unit_price":          prec,
                "line_total":          tot_partida,
                "following_document":  doc_sig,
                "line_number":         line_number,
            })

    wb.close()

    return {
        "orders":     orders,
        "raw_rows":   raw_rows,
        "header_row": header_row_idx,
        "total_rows": total_data_rows,
    }
